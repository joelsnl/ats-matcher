from __future__ import annotations

import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

import ftfy

LOW_TEXT_THRESHOLD = 200
PAGE_MARKER = "--- page {n} ---"
Kind = Literal["pdf", "docx", "txt"]


class UnsupportedDocumentError(ValueError):
    pass


@dataclass
class ExtractedDocument:
    text: str
    pages: list[str] = field(default_factory=list)
    quality: Literal["ok", "low"] = "ok"
    source: str = "unknown"
    char_count: int = 0
    path: str | None = None

    def __post_init__(self) -> None:
        self.char_count = len(self.text)
        if self.char_count < LOW_TEXT_THRESHOLD:
            self.quality = "low"


def detect_kind(path: Path) -> Kind:
    suffix = path.suffix.lower()
    magic = path.read_bytes()[:8]
    if magic.startswith(b"%PDF") or suffix == ".pdf":
        return "pdf"
    if magic.startswith(b"PK") and suffix in {".docx", ".doc"}:
        return "docx"
    if suffix in {".docx", ".doc"}:
        return "docx"
    if suffix in {".txt", ".md"}:
        return "txt"
    raise UnsupportedDocumentError(f"Unsupported document type: {path.suffix or 'unknown'}")


def extract_document(path: str | Path) -> ExtractedDocument:
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(path)
    kind = detect_kind(path)
    if kind == "pdf":
        doc = _extract_pdf(path)
    elif kind == "docx":
        doc = _extract_docx(path)
    else:
        raw = path.read_text(encoding="utf-8", errors="replace")
        doc = ExtractedDocument(text=_normalize(raw), pages=[_normalize(raw)], source="txt")
    doc.path = str(path)
    doc.text = _with_page_markers(doc.pages) if len(doc.pages) > 1 else doc.text
    doc.char_count = len(doc.text)
    if doc.char_count < LOW_TEXT_THRESHOLD:
        doc.quality = "low"
    return doc


def _extract_pdf(path: Path) -> ExtractedDocument:
    pages, source = _pdfplumber_pages(path)
    joined = "\n\n".join(p for p in pages if p.strip())
    if len(joined.strip()) < LOW_TEXT_THRESHOLD:
        fallback, fb_source = _pypdf_pages(path)
        fb_joined = "\n\n".join(p for p in fallback if p.strip())
        if len(fb_joined) > len(joined):
            pages, source, joined = fallback, fb_source, fb_joined
    pages = [_normalize(p) for p in pages]
    pages = _drop_repeated_headers(pages)
    text = _with_page_markers(pages) if len(pages) > 1 else (pages[0] if pages else "")
    return ExtractedDocument(text=text, pages=pages, source=source)


def _pdfplumber_pages(path: Path) -> tuple[list[str], str]:
    import pdfplumber

    pages: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            chunks: list[str] = []
            layout = page.extract_text(layout=True) or ""
            plain = page.extract_text() or ""
            chosen = _prefer_two_column(page, layout, plain)
            if chosen:
                chunks.append(chosen)
            for table in page.extract_tables() or []:
                rows = [" | ".join((cell or "").strip() for cell in row) for row in table]
                table_text = "\n".join(r for r in rows if r.strip(" |"))
                if table_text and table_text not in chosen:
                    chunks.append(table_text)
            pages.append("\n".join(chunks).strip())
    return pages, "pdfplumber"


def _prefer_two_column(page: object, layout: str, plain: str) -> str:
    """If the page looks two-column, reconstruct left-then-right reading order."""
    words = getattr(page, "extract_words", lambda **_: [])() or []
    if len(words) < 12:
        return layout or plain
    width = float(getattr(page, "width", 0) or 0)
    if width <= 0:
        return layout or plain
    mid = width / 2
    left = [w for w in words if float(w.get("x0", 0)) < mid - 8]
    right = [w for w in words if float(w.get("x0", 0)) >= mid - 8]
    if not left or not right:
        return layout or plain
    # Need a real gap between columns, not just a wide single column.
    left_max = max(float(w.get("x1", 0)) for w in left)
    right_min = min(float(w.get("x0", 0)) for w in right)
    if right_min - left_max < 12:
        return layout or plain
    if len(left) < 4 or len(right) < 4:
        return layout or plain
    return _words_to_text(left) + "\n\n" + _words_to_text(right)


def _words_to_text(words: list[dict]) -> str:
    lines: list[list[dict]] = []
    for word in sorted(words, key=lambda w: (round(float(w.get("top", 0)), 0), float(w.get("x0", 0)))):
        if not lines:
            lines.append([word])
            continue
        prev_top = float(lines[-1][0].get("top", 0))
        if abs(float(word.get("top", 0)) - prev_top) < 4:
            lines[-1].append(word)
        else:
            lines.append([word])
    return "\n".join(" ".join(w.get("text", "") for w in line) for line in lines)


def _pypdf_pages(path: Path) -> tuple[list[str], str]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    return pages, "pypdf"


def _extract_docx(path: Path) -> ExtractedDocument:
    from docx import Document

    document = Document(str(path))
    parts: list[str] = []
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = _normalize("\n".join(parts))
    return ExtractedDocument(text=text, pages=[text], source="docx")


def _with_page_markers(pages: list[str]) -> str:
    chunks = []
    for i, page in enumerate(pages, start=1):
        chunks.append(f"{PAGE_MARKER.format(n=i)}\n{page}".rstrip())
    return "\n\n".join(chunks)


def _drop_repeated_headers(pages: list[str]) -> list[str]:
    if len(pages) < 2:
        return pages
    first_lines = []
    last_lines = []
    for page in pages:
        lines = [ln.strip() for ln in page.splitlines() if ln.strip()]
        if lines:
            first_lines.append(lines[0])
            last_lines.append(lines[-1])
    drop: set[str] = set()
    for bucket in (first_lines, last_lines):
        if not bucket:
            continue
        value, count = Counter(bucket).most_common(1)[0]
        if count == len(pages) and 0 < len(value) < 80:
            drop.add(value)

    cleaned = []
    for page in pages:
        kept = [ln for ln in page.splitlines() if ln.strip() not in drop]
        cleaned.append("\n".join(kept).strip())
    return cleaned


_MULTI_NL = re.compile(r"\n{3,}")
_MULTI_SPACE = re.compile(r"[ \t]{2,}")


def _normalize(text: str) -> str:
    text = ftfy.fix_text(text or "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _MULTI_SPACE.sub(" ", text)
    text = _MULTI_NL.sub("\n\n", text)
    return text.strip()
