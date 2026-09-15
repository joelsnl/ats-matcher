from __future__ import annotations

import http.client
import json
import threading
from pathlib import Path

import pytest

from ats_matcher.llm.engine import LlamaEngine
from ats_matcher.server import MAX_UPLOAD, create_server


MINIMAL_EXTRACT = {
    "full_name": "Taylor Example",
    "email": None,
    "skills": ["React", "SQL"],
    "years_experience": 5,
    "primary_industry": "software",
    "location": {"city": "Berlin", "region": None, "country": "Germany", "raw": "Berlin"},
    "github_url": None,
    "linkedin_url": None,
    "recent_titles": ["Frontend Developer"],
    "employers": [{"name": "Example Co"}],
    "certifications": [],
    "spoken_languages": [{"name": "English", "level": "fluent"}],
    "birth_year": None,
    "age_estimate": None,
    "age_source": None,
    "confidence": 0.8,
}


def dummy_engine(extract=None, letter=None):
    payload = extract or MINIMAL_EXTRACT
    note = letter or (
        "Hi Example team,\n\n"
        "At Example Co I spent five years building React and SQL services for a product people used every day. "
        "The listing asks for that same mix of interface work and data, which is the work I already do.\n\n"
        "I am in Berlin and looking for a hybrid team. Happy to talk through a recent project if useful.\n\n"
        "Taylor Example\n"
    )

    def complete(messages):
        joined = " ".join(item["content"] for item in messages)
        if "plain-text application note" in joined:
            return note
        if "interview stories for one real job" in joined:
            return json.dumps({
                "briefing": "This role will probe React work you already did at Example Co, plus how you explain data trade-offs.",
                "stories": [
                    {"prompt": "Walk through a time you had to explain a data trade-off to someone outside engineering.", "anchor": "Example Co, SQL reporting"},
                    {"prompt": "Tell me about a production issue you helped contain.", "anchor": "Example Co, React services"},
                    {"prompt": "When did you disagree with a teammate about an approach, and what changed?", "anchor": "Example Co"},
                ],
                "ask_them": ["What does a useful first month look like on this team?"],
            })
        if "infer market-standard job titles" in joined:
            return json.dumps({"possible_titles": ["Frontend Engineer", "Software Engineer"]})
        return json.dumps(payload)

    return LlamaEngine(complete_fn=complete)


def serve(engine=None, mode="cpu", translator=None):
    server = create_server(0, engine=engine, mode=mode, translator=translator)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    return server, thread


@pytest.fixture
def app_server():
    server, thread = serve()
    yield server
    server.shutdown()
    server.server_close()
    thread.join(timeout=3)


@pytest.fixture
def model_server():
    server, thread = serve(dummy_engine())
    yield server
    server.shutdown()
    server.server_close()
    thread.join(timeout=3)


@pytest.fixture
def basic_server():
    server, thread = serve(mode="basic")
    yield server
    server.shutdown()
    server.server_close()
    thread.join(timeout=3)


def request(server, method, route, body=None, headers=None):
    connection = http.client.HTTPConnection("127.0.0.1", server.server_port, timeout=5)
    try:
        connection.request(method, route, body=body, headers=headers or {})
        response = connection.getresponse()
        return response.status, response.read(), dict(response.getheaders())
    finally:
        connection.close()


def test_only_public_app_assets_are_served(app_server):
    status, body, headers = request(app_server, "GET", "/")
    assert status == 200
    assert b"ATS Matcher" in body
    assert b"Alex Morgan" not in body
    assert b"About this prototype" not in body
    assert b"Prototype" not in body
    assert headers["Cache-Control"] == "no-store"
    for route in ("/../cv.json", "/cv.json", "/.env", "/%2e%2e/resume.pdf"):
        assert request(app_server, "GET", route)[0] == 404


def test_parse_uses_local_model(model_server):
    body = b"Taylor Example\ntaylor@example.com\nReact and SQL\n5 years experience\n"
    status, payload, _ = request(model_server, "POST", "/api/parse", body, {"X-CV-Extension": "txt"})
    data = json.loads(payload)
    assert status == 200
    assert data["mode"] == "local-model"
    assert "React and SQL" in data["text"]
    assert data["profile"]["full_name"] == "Taylor Example"
    assert data["profile"]["possible_titles"] == ["Frontend Engineer", "Software Engineer"]
    assert "email" not in data["profile"]
    assert "birth_year" not in data["profile"]


def test_parse_without_model_is_unavailable(app_server):
    status, payload, _ = request(app_server, "POST", "/api/parse", b"Taylor Example", {"X-CV-Extension": "txt"})
    assert status == 503
    assert "GGUF" in json.loads(payload)["error"]


def test_basic_mode_reads_cv_without_model(basic_server):
    body = b"Taylor Example\ntaylor@example.com\nReact and SQL\n5 years experience\n"
    status, payload, _ = request(basic_server, "GET", "/api/health")
    assert status == 200
    health = json.loads(payload)
    assert health["mode"] == "basic"
    assert health["device"] is None
    status, payload, _ = request(basic_server, "POST", "/api/parse", body, {"X-CV-Extension": "txt"})
    data = json.loads(payload)
    assert status == 200
    assert data["mode"] == "basic"
    assert "React and SQL" in data["text"]
    assert "email" not in data["profile"]


def test_foreign_origin_and_host_are_rejected(app_server):
    for headers in ({"Origin": "https://unrelated.example"}, {"Host": "unrelated.example"}):
        assert request(app_server, "GET", "/", headers=headers)[0] == 403
    headers = {"Origin": f"http://127.0.0.1:{app_server.server_port}"}
    assert request(app_server, "GET", "/api/health", headers=headers)[0] == 200


@pytest.mark.parametrize("header,value", [
    ("Host", "localhost.attacker.example:8765"),
    ("Host", "127.0.0.1.attacker.example:8765"),
    ("Host", "localhost:1"),
    ("Origin", "http://localhost.attacker.example"),
    ("Origin", "http://127.0.0.1.attacker.example"),
    ("Origin", "http://localhost:1"),
])
def test_localhost_prefixes_do_not_bypass_origin_checks(app_server, header, value):
    status, _, headers = request(app_server, "GET", "/", headers={header: value})
    assert status == 403
    assert headers.get("Access-Control-Allow-Origin") != value


@pytest.mark.parametrize("error", [ConnectionAbortedError, ConnectionResetError, BrokenPipeError])
def test_browser_disconnect_is_handled(app_server, monkeypatch, error):
    from http.server import BaseHTTPRequestHandler

    def disconnect(self):
        raise error("browser disconnected")

    handler = object.__new__(app_server.RequestHandlerClass)
    monkeypatch.setattr(BaseHTTPRequestHandler, "handle", disconnect)
    handler.handle()
    assert handler.close_connection


def test_unexpected_handler_errors_are_not_hidden(app_server, monkeypatch):
    from http.server import BaseHTTPRequestHandler

    def fail(self):
        raise RuntimeError("unexpected failure")

    handler = object.__new__(app_server.RequestHandlerClass)
    monkeypatch.setattr(BaseHTTPRequestHandler, "handle", fail)
    with pytest.raises(RuntimeError, match="unexpected failure"):
        handler.handle()


def test_upload_limits_and_file_types(app_server):
    assert request(app_server, "POST", "/api/parse", b"x", {"X-CV-Extension": "exe"})[0] == 400
    assert request(app_server, "POST", "/api/parse", b"", {"X-CV-Extension": "pdf"})[0] == 400
    assert request(app_server, "POST", "/api/parse", headers={"X-CV-Extension": "pdf", "Content-Length": str(MAX_UPLOAD + 1)})[0] == 413


def test_failed_parse_cleans_up_upload_and_does_not_leak_details(model_server, monkeypatch):
    observed = []

    def fail(cv_path, **kwargs):
        observed.append(Path(cv_path))
        assert Path(cv_path).exists()
        raise ValueError("private CV content")

    monkeypatch.setattr("ats_matcher.server.parse_cv", fail)
    status, body, _ = request(model_server, "POST", "/api/parse", b"bad pdf", {"X-CV-Extension": "pdf"})
    assert status == 422
    assert b"private CV content" not in body
    assert observed and not observed[0].exists()


@pytest.mark.parametrize("extension", ["pdf", "docx"])
def test_binary_cv_upload_reads_real_document(model_server, extension):
    from io import BytesIO

    stream = BytesIO()
    if extension == "pdf":
        from reportlab.pdfgen import canvas

        document = canvas.Canvas(stream)
        document.drawString(72, 750, "Taylor Example")
        document.drawString(72, 730, "React, TypeScript, SQL")
        document.save()
    else:
        from docx import Document

        document = Document()
        document.add_paragraph("Taylor Example")
        document.add_paragraph("React, TypeScript, SQL")
        document.save(stream)
    status, payload, _ = request(
        model_server, "POST", "/api/parse", stream.getvalue(),
        {"X-CV-Extension": extension},
    )
    assert status == 200
    assert "Taylor Example" in json.loads(payload)["text"]
    assert "TypeScript" in json.loads(payload)["text"]


def test_search_api_uses_common_query_and_registry(app_server, monkeypatch):
    from ats_matcher.schemas.jobs import JobSearchResponse, JobsMeta, Job
    observed = []
    def search(self, query, providers=None, **kwargs):
        observed.append((query, providers))
        return JobSearchResponse(jobs=[Job(title="Engineer", company="Example", job_id="123", application_url="https://www.linkedin.com/jobs/view/123")], jobs_meta=JobsMeta(requested=query.limit, returned=1, provider="linkedin", query_used=query))
    monkeypatch.setattr("ats_matcher.server.JobSearchService.search", search)
    status, body, _ = request(app_server, "POST", "/api/jobs", json.dumps({"keyword": "engineer", "remoteFilter": "remote", "providers": ["linkedin"]}), {"Content-Type": "application/json"})
    assert status == 200
    assert json.loads(body)["jobs"][0]["title"] == "Engineer"
    assert observed[0][0].workplace_type == "remote"
    assert observed[0][1] == ["linkedin"]
    status, body, _ = request(app_server, "GET", "/api/providers")
    assert status == 200 and json.loads(body)["providers"][0]["id"] == "linkedin"


def test_search_api_rejects_invalid_query(app_server):
    for body in ('{"keyword":"engineer","limit":0}', '{"keyword":"engineer","providers":"linkedin"}', '[]', '{bad'):
        assert request(app_server, "POST", "/api/jobs", body)[0] == 400
    assert request(app_server, "POST", "/api/jobs", headers={"Content-Length": "20000"})[0] == 400


@pytest.mark.parametrize("outcome,http_status", [("error", 502), ("partial", 200)])
def test_search_api_distinguishes_failed_and_partial_sources(app_server, monkeypatch, outcome, http_status):
    from ats_matcher.schemas.jobs import JobSearchResponse, JobsMeta
    def search(self, query, providers=None, **kwargs):
        return JobSearchResponse(jobs_meta=JobsMeta(requested=query.limit, returned=0, provider="linkedin", status=outcome, errors=["linkedin: Unavailable"]))
    monkeypatch.setattr("ats_matcher.server.JobSearchService.search", search)
    status, body, _ = request(app_server, "POST", "/api/jobs", '{"keyword":"engineer"}')
    assert status == http_status
    assert json.loads(body)["jobs_meta"]["status"] == outcome


def test_search_api_multiple_filters_and_custom_days(app_server, monkeypatch):
    from ats_matcher.schemas.jobs import JobSearchResponse, JobsMeta
    observed = []
    def search(self, query, providers=None, **kwargs):
        observed.append(query)
        return JobSearchResponse(jobs_meta=JobsMeta(requested=query.limit, returned=0, provider="linkedin", query_used=query))
    monkeypatch.setattr("ats_matcher.server.JobSearchService.search", search)
    body = {"keyword": "engineer", "jobType": ["part time", "internship"], "remoteFilter": ["remote", "hybrid"], "experienceLevel": ["entry level", "associate"], "postedWithinDays": 3}
    status, payload, _ = request(app_server, "POST", "/api/jobs", json.dumps(body), {"Content-Type": "application/json"})
    assert status == 200
    assert json.loads(payload)["jobs_meta"]["query_used"]["posted_within_days"] == 3
    assert set(observed[0].workplace_type) == {"remote", "hybrid"}
    for bad in ({"remoteFilter": []}, {"postedWithinDays": -1}, {"jobType": ["made up"]}):
        assert request(app_server, "POST", "/api/jobs", json.dumps({**body, **bad}))[0] == 400
    assert len(observed) == 1


@pytest.mark.parametrize("value,expected", [
    (["on_site"], ["on_site"]),
    (["On-Site"], ["on_site"]),
    (["ON SITE", "Remote", "HYBRID"], ["hybrid", "on_site", "remote"]),
    ("On-Site", "on_site"),
])
def test_search_api_workplace_shapes_and_capitalization(app_server, monkeypatch, value, expected):
    from ats_matcher.schemas.jobs import JobSearchResponse, JobsMeta
    observed = []
    def search(self, query, providers=None, **kwargs):
        observed.append(query)
        return JobSearchResponse(jobs_meta=JobsMeta(requested=query.limit, returned=0, provider="linkedin", query_used=query))
    monkeypatch.setattr("ats_matcher.server.JobSearchService.search", search)
    status, payload, _ = request(app_server, "POST", "/api/jobs", json.dumps({"keywords": "engineer", "workplace_type": value}))
    assert status == 200
    assert observed[0].workplace_type == expected
    assert json.loads(payload)["jobs_meta"]["query_used"]["workplace_type"] == expected


def test_search_validation_error_is_readable(app_server):
    status, payload, _ = request(app_server, "POST", "/api/jobs", json.dumps({"keywords": "engineer", "workplace_type": ["INVALID_PRIVATE_INPUT"]}))
    assert status == 400
    message = json.loads(payload)["error"]
    assert "work arrangement" in message and "On-site" in message
    assert all(value not in message for value in ["literal_error", "pydantic", "JobSearchQuery", "INVALID_PRIVATE_INPUT"])


def test_live_workspace_assets_and_health(app_server):
    status, payload, _ = request(app_server, "GET", "/api/health")
    assert status == 200
    assert json.loads(payload)["sample_jobs"] is False
    assert json.loads(payload)["live_jobs"] is True
    assert json.loads(payload)["translate"] is True
    assert json.loads(payload)["languages"][0]["id"] == "en"
    assert json.loads(payload)["mode"] == "unavailable"
    assert json.loads(payload)["device"] is None
    assert json.loads(payload)["cover_letter"] is False
    assert json.loads(payload)["interview_prep"] is True
    assert json.loads(payload)["interview_stories"] is False
    status, payload, headers = request(app_server, "GET", "/jobs-store.js")
    assert status == 200 and "javascript" in headers["Content-Type"]
    status, payload, headers = request(app_server, "GET", "/theme.js")
    assert status == 200 and "javascript" in headers["Content-Type"]


def test_translate_api_keeps_original_text_and_language():
    from ats_matcher.jobs.translate import Translation

    class FakeTranslator:
        def translate(self, text, target="en", source="auto"):
            assert "Wij zoeken" in text
            assert target == "en"
            return Translation("We are hiring an engineer for Azure.", "nl", "en", True)

    server, thread = serve(translator=FakeTranslator())
    try:
        status, payload, _ = request(server, "POST", "/api/translate", json.dumps({"text": "Wij zoeken een engineer.", "target": "en"}), {"Content-Type": "application/json"})
        data = json.loads(payload)
        assert status == 200
        assert data["translated"] is True
        assert data["source_language"] == "nl"
        assert "Azure" in data["text"]
        assert request(server, "POST", "/api/translate", json.dumps({"text": "", "target": "en"}), {"Content-Type": "application/json"})[0] == 400
        assert request(server, "POST", "/api/translate", json.dumps({"text": "Hallo", "target": "xx"}), {"Content-Type": "application/json"})[0] == 400
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=3)


def test_resolve_model_matches_cli_defaults(tmp_path):
    from ats_matcher.config import Settings
    from ats_matcher.server import resolve_model

    gguf = tmp_path / "model.gguf"
    other = tmp_path / "other.gguf"
    settings = Settings.model_construct(llama_model_path=gguf)
    assert resolve_model(settings=settings, search_root=tmp_path) == gguf
    assert resolve_model(other, settings=settings, search_root=tmp_path) == other
    assert resolve_model(settings=Settings.model_construct(llama_model_path=None), search_root=tmp_path) is None


def test_resolve_model_discovers_gguf_in_models_dir(tmp_path):
    from ats_matcher.config import Settings
    from ats_matcher.server import resolve_model

    folder = tmp_path / ".models"
    folder.mkdir()
    (folder / "weaker-Q3_K_M.gguf").write_bytes(b"x")
    preferred = folder / "Qwen2.5-7B-Instruct-Q4_K_M.gguf"
    preferred.write_bytes(b"x")
    settings = Settings.model_construct(llama_model_path=None)
    assert resolve_model(settings=settings, search_root=tmp_path) == preferred


def test_gpu_layers_follow_mode():
    from ats_matcher.config import Settings
    from ats_matcher.server import gpu_layers_for_mode

    cpu_settings = Settings.model_construct(llama_n_gpu_layers=-1)
    assert gpu_layers_for_mode("cpu", cpu_settings) == 0
    assert gpu_layers_for_mode("basic", cpu_settings) == 0
    assert gpu_layers_for_mode("gpu", Settings.model_construct(llama_n_gpu_layers=0)) == -1
    assert gpu_layers_for_mode("gpu", Settings.model_construct(llama_n_gpu_layers=20)) == 20


def test_gpu_mode_health_reports_device():
    server, thread = serve(dummy_engine(), mode="gpu")
    try:
        status, payload, _ = request(server, "GET", "/api/health")
        assert status == 200
        data = json.loads(payload)
        assert data["mode"] == "local-model"
        assert data["device"] == "gpu"
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=3)


def test_parse_uses_injected_llm_engine_by_default_path():
    server, thread = serve(dummy_engine())
    try:
        status, payload, _ = request(server, "GET", "/api/health")
        assert status == 200
        assert json.loads(payload)["mode"] == "local-model"
        assert json.loads(payload)["device"] == "cpu"
        body = b"Taylor Example\nReact, SQL\nFrontend Developer at Example Co\n"
        status, payload, _ = request(server, "POST", "/api/parse", body, {"X-CV-Extension": "txt"})
        data = json.loads(payload)
        assert status == 200
        assert data["mode"] == "local-model"
        assert data["profile"]["full_name"] == "Taylor Example"
        assert data["profile"]["possible_titles"] == ["Frontend Engineer", "Software Engineer"]
        assert data["profile"]["skills"] == ["React", "SQL"]
        assert data["profile"]["years_experience"] == 5
        assert data["profile"]["employers"][0]["name"] == "Example Co"
        assert "email" not in data["profile"]
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=3)


def test_cover_letter_api_writes_from_profile_and_listing():
    server, thread = serve(dummy_engine())
    try:
        health = json.loads(request(server, "GET", "/api/health")[1])
        assert health["cover_letter"] is True
        assert health["interview_prep"] is True
        assert health["interview_stories"] is True
        body = {
            "profile": {
                "name": "Taylor Example",
                "skills": ["React", "SQL"],
                "location": "Berlin, Germany",
                "years_experience": 5,
                "employers": [{"name": "Example Co"}],
            },
            "job": {
                "title": "Frontend Engineer",
                "company": "Real Source Example",
                "description": "We need React, SQL and communication for a product team in Berlin.",
            },
        }
        status, payload, _ = request(server, "POST", "/api/cover-letter", json.dumps(body), {"Content-Type": "application/json"})
        data = json.loads(payload)
        assert status == 200
        assert "Example Co" in data["letter"] or "React" in data["letter"]
        assert "I am writing to" not in data["letter"].lower()
        assert request(server, "POST", "/api/cover-letter", json.dumps({"profile": {"name": "Taylor"}, "job": {"title": "A", "company": "B", "description": "short"}}), {"Content-Type": "application/json"})[0] == 400
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=3)


def test_cover_letter_requires_local_model(basic_server):
    status, payload, _ = request(basic_server, "POST", "/api/cover-letter", json.dumps({
        "profile": {"name": "Taylor Example", "skills": ["React"]},
        "job": {"title": "Engineer", "company": "Example", "description": "A long enough job description about React work."},
    }), {"Content-Type": "application/json"})
    assert status == 503
    assert "local model" in json.loads(payload)["error"].lower()


def test_interview_prep_api_matches_open_materials_and_stories():
    server, thread = serve(dummy_engine())
    try:
        body = {
            "profile": {
                "name": "Taylor Example",
                "skills": ["React", "SQL", "Jenkins"],
                "roles": ["Platform Engineer"],
                "employers": [{"name": "Example Co"}],
            },
            "job": {
                "title": "Platform Engineer",
                "company": "Real Source Example",
                "description": "We need Python, Jenkins, Kubernetes and CI/CD for a platform team in Berlin.",
            },
        }
        status, payload, _ = request(server, "POST", "/api/interview-prep", json.dumps(body), {"Content-Type": "application/json"})
        data = json.loads(payload)
        assert status == 200
        assert data["generated"] is True
        assert "Example Co" in data["briefing"]
        assert data["stories"]
        assert data["stories"][0]["prompt"]
        titles = [item["title"] for item in data["resources"]]
        assert "Behavioral interview guide" in titles
        assert any("DevOps" in title or "Jenkins" in title for title in titles)
        assert all(item["url"].startswith("https://") for item in data["resources"])
        assert request(server, "POST", "/api/interview-prep", json.dumps({"profile": {"name": "Taylor"}, "job": {"title": "A", "company": "B", "description": "short"}}), {"Content-Type": "application/json"})[0] == 400
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=3)


def test_interview_prep_works_without_a_model(basic_server):
    status, payload, _ = request(basic_server, "POST", "/api/interview-prep", json.dumps({
        "profile": {"name": "Taylor Example", "skills": ["React"]},
        "job": {"title": "Frontend Engineer", "company": "Example", "description": "A long enough job description about React and TypeScript work."},
    }), {"Content-Type": "application/json"})
    data = json.loads(payload)
    assert status == 200
    assert data["generated"] is False
    assert data["resources"]
    assert data["briefing"]
    assert data["ask_them"]
    titles = [item["title"] for item in data["resources"]]
    assert "Behavioral interview guide" in titles
    assert "Kubernetes the Hard Way" not in titles


def test_prototype_module_is_server_alias():
    import ats_matcher.prototype as proto
    import ats_matcher.server as server

    assert proto.create_server is server.create_server
    assert proto.ASSETS == server.ASSETS
    assert proto.ASSETS.name == "app"


def coaching_request():
    return {"profile": {"name": "Taylor", "skills": ["Python"], "achievements": "I wrote Python import checks."},
            "job": {"title": "Engineer", "company": "Example", "description": "Required: Python and Kubernetes. Communication with stakeholders is important."}}


def test_coaching_assets_and_analysis_work_without_a_model(basic_server):
    assert request(basic_server, "GET", "/coaching.js")[0] == 200
    assert request(basic_server, "GET", "/desk.js")[0] == 200
    assert request(basic_server, "GET", "/desk-logic.js")[0] == 200
    status, body, _ = request(basic_server, "POST", "/api/role-analysis", json.dumps(coaching_request()))
    assert status == 200
    rows = {r["name"]: r for r in json.loads(body)["requirements"]}
    assert rows["Python"]["profile_evidence"] == "I wrote Python import checks."
    assert rows["Kubernetes"]["status"] == "not_evidenced"
    assert rows["Communication"]["kind"] == "soft"


def test_practice_api_returns_a_validated_single_session(basic_server):
    data = coaching_request()
    status, body, _ = request(basic_server, "POST", "/api/interview-prep", json.dumps(data))
    assert status == 200
    lesson = next(l for l in json.loads(body)["learning_plan"]["lessons"] if l["skill"] == "Kubernetes")
    status, body, _ = request(basic_server, "POST", "/api/practice-lesson", json.dumps({**data, "skill_id": lesson["id"], "mode": "refresher"}))
    assert status == 200
    result = json.loads(body)["lesson"]
    assert result["mode"] == "refresher" and result["status"] == "not_evidenced"
    assert result["criteria"] and result["exercise"]
    assert request(basic_server, "POST", "/api/practice-lesson", json.dumps({**data, "skill_id": "forged"}))[0] == 400
    assert request(basic_server, "POST", "/api/practice-lesson", json.dumps({**data, "skill_id": lesson["id"], "mode": "expert"}))[0] == 400


def test_manual_letter_review_needs_no_model(basic_server):
    data = {**coaching_request(), "letter": "I operated AWS systems for 42 years and wrote Python tools. My experience is useful for the work described by your platform team.\n\nTaylor"}
    status, body, _ = request(basic_server, "POST", "/api/cover-letter-review", json.dumps(data))
    assert status == 200
    review = json.loads(body)
    assert any("AWS" in issue for issue in review["issues"])
    assert any("42" in issue for issue in review["issues"])
    assert review["context"]["requirements"]


@pytest.mark.parametrize("route", ["/api/role-analysis", "/api/practice-lesson", "/api/cover-letter-review", "/api/practice-feedback", "/api/cover-letter-revise"])
def test_coaching_routes_reject_bad_inputs_and_foreign_origins(basic_server, route):
    for body in ("[]", "null", '{bad', '{"profile": [], "job": {}}'):
        assert request(basic_server, "POST", route, body)[0] == 400
    assert request(basic_server, "POST", route, headers={"Content-Length": "96001"})[0] == 400
    assert request(basic_server, "POST", route, json.dumps(coaching_request()), {"Origin": "https://unrelated.example"})[0] == 403


def test_letter_options_and_review_are_returned(model_server):
    data = {**coaching_request(), "options": {"tone": "warm", "length": "concise", "language": "nl"}}
    # Reject invalid options before invoking the model.
    assert request(model_server, "POST", "/api/cover-letter", json.dumps({**data, "options": {"tone": "invent achievements"}}))[0] == 400
    # This engine is deterministic; the assertion checks plumbing, not language quality.
    status, body, _ = request(model_server, "POST", "/api/cover-letter", json.dumps({**data, "profile": {**data["profile"], "skills": ["React", "SQL"]}}))
    assert status == 200
    result = json.loads(body)
    assert result["options"] == data["options"]
    assert result["review"]["word_count"] > 0


def test_practice_feedback_api_validates_the_attempt_and_exercise(basic_server):
    data = coaching_request()
    pack = json.loads(request(basic_server, "POST", "/api/interview-prep", json.dumps(data))[1])
    lesson = pack["learning_plan"]["lessons"][0]
    payload = {**data, "skill_id": lesson["id"], "mode": lesson["mode"], "lesson": lesson,
               "answer": "I would inspect the Service selector and Pod labels first, then check readiness."}
    status, body, _ = request(basic_server, "POST", "/api/practice-feedback", json.dumps(payload))
    assert status == 200
    assert json.loads(body)["generated"] is False
    for invalid in ({"answer": "ok"}, {"lesson": None}, {"lesson": {"exercise": "Too short", "criteria": []}}):
        assert request(basic_server, "POST", "/api/practice-feedback", json.dumps({**payload, **invalid}))[0] == 400


def test_passage_revision_requires_a_model(basic_server):
    status, body, _ = request(basic_server, "POST", "/api/cover-letter-revise", json.dumps(coaching_request()))
    assert status == 503
    assert "original text is kept" in json.loads(body)["error"]


def test_passage_revision_api_validation_and_lock_release():
    replacement = "I wrote Python import checks and can explain how I approached that work."
    server, thread = serve(LlamaEngine(complete_fn=lambda messages: replacement))
    try:
        passage = "I wrote Python import checks and would like to explain them to your team."
        letter = "Dear Example team,\n\n" + passage + "\n\nThe validation work in this role interests me. I would welcome the opportunity to discuss how these checks connect with the requirements of this position and learn about your team.\n\nTaylor"
        payload = {**coaching_request(), "letter": letter, "start": letter.index(passage),
                   "end": letter.index(passage) + len(passage), "instruction": "plain"}
        # A rejected request must release the model lock for the next valid request.
        assert request(server, "POST", "/api/cover-letter-revise", json.dumps({**payload, "start": -1}))[0] == 400
        status, body, _ = request(server, "POST", "/api/cover-letter-revise", json.dumps(payload))
        assert status == 200
        assert json.loads(body)["replacement"] == replacement
    finally:
        server.shutdown()
        server.server_close()
        thread.join(timeout=3)
