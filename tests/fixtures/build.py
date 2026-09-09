from __future__ import annotations

from pathlib import Path


def build_all(dest: Path) -> Path:
    dest.mkdir(parents=True, exist_ok=True)
    write_digital_pdf(dest / "digital.pdf")
    write_two_column_pdf(dest / "two_column.pdf")
    write_multipage_pdf(dest / "multipage.pdf")
    write_low_text_pdf(dest / "low_text.pdf")
    write_docx(dest / "sample.docx")
    return dest


def write_digital_pdf(path: Path) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    y = height - 56
    lines = [
        "Maya Chen",
        "Backend Engineer",
        "Berlin, Germany",
        "maya.chen@example.com  |  +49 30 12345678",
        "GitHub: https://github.com/maya-chen  LinkedIn: https://www.linkedin.com/in/mayachen",
        "",
        "SUMMARY",
        "Software engineer with 6 years building APIs in Python and Go.",
        "",
        "SKILLS",
        "Python, FastAPI, PostgreSQL, Docker, Kubernetes, Redis",
        "",
        "EXPERIENCE",
        "Senior Backend Engineer, Nordwind Labs  2021 - Present",
        "Backend Engineer, Helio Systems  2018 - 2021",
        "",
        "EDUCATION",
        "B.Sc. Computer Science, TU Berlin  2014 - 2018",
    ]
    for line in lines:
        c.drawString(50, y, line)
        y -= 16
    c.save()
    return path


def write_two_column_pdf(path: Path) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    y = height - 60
    left = [
        "LEFTUNIQUE",
        "Omar Haddad",
        "omar.haddad@example.net",
        "https://github.com/ohaddad",
        "Skills",
        "TypeScript",
        "React",
        "Node.js",
    ]
    right = [
        "RIGHTUNIQUE",
        "Experience",
        "Frontend Engineer, 2020-2024",
        "Built design systems",
        "Location: Amsterdam, NL",
        "Industry: software",
    ]
    for line in left:
        c.drawString(48, y, line)
        y -= 18
    y = height - 60
    for line in right:
        c.drawString(320, y, line)
        y -= 18
    c.save()
    return path


def write_multipage_pdf(path: Path) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    width, height = A4
    c.drawString(50, height - 40, "HeaderRepeat")
    c.drawString(50, height - 80, "PageOneBody ALPHA")
    c.drawString(50, 40, "FooterRepeat")
    c.showPage()
    c.drawString(50, height - 40, "HeaderRepeat")
    c.drawString(50, height - 80, "PageTwoBody BETA")
    c.drawString(50, 40, "FooterRepeat")
    c.save()
    return path


def write_low_text_pdf(path: Path) -> Path:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=A4)
    c.drawString(80, 500, "Hi")
    c.save()
    return path


def write_docx(path: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("Priya Nair", level=1)
    doc.add_paragraph("Data Scientist  |  London, United Kingdom")
    doc.add_paragraph("priya.nair@example.com")
    doc.add_paragraph("https://github.com/priyanair")
    doc.add_paragraph("Skills: Python, pandas, scikit-learn, SQL, dbt")
    doc.add_paragraph("Experience: Data Scientist at Maple Analytics, 2019 - Present")
    doc.save(str(path))
    return path
